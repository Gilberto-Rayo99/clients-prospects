"""Export de prospectos a Excel con dropdown de estado y colores."""
from __future__ import annotations

import logging
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import config

logger = logging.getLogger(__name__)

# Colores
FILL_GREEN = PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid")
FILL_YELLOW = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
FILL_HEADER = PatternFill(start_color="1F6FEB", end_color="1F6FEB", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

COLUMNS = [
    ("Score", 8),
    ("Nombre", 32),
    ("Categoría", 22),
    ("Dirección", 42),
    ("Teléfono", 18),
    ("Email", 30),
    ("Canal de contacto", 18),
    ("Valor canal", 38),
    ("Estrellas", 10),
    ("Reseñas", 10),
    ("Tiene web", 12),
    ("Link Maps", 22),
    ("Link Landing", 28),
    ("Estado", 22),
]


def _channel_label(channel: str) -> str:
    return {
        "email": "Email",
        "whatsapp": "WhatsApp",
        "facebook": "Facebook",
        "visit": "Visita presencial",
    }.get(channel, channel)


def export_to_excel(prospects: list[dict]) -> str:
    """Genera Excel y devuelve la ruta. Filas coloreadas por score."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Prospectos"

    # Header
    for col_idx, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=name)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = FILL_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"

    # Filas
    for row_idx, p in enumerate(prospects, start=2):
        score = p.get("score", 0)
        values = [
            score,
            p.get("name", ""),
            p.get("category", ""),
            p.get("address", ""),
            p.get("phone", "") or "",
            p.get("email", "") or "",
            _channel_label(p.get("contact_channel", "")),
            p.get("contact_value", "") or "",
            p.get("rating") or "",
            p.get("reviews_count") or 0,
            "Sí" if p.get("website") else "No",
            p.get("maps_url") or "",
            p.get("landing_path") or "",
            "Pendiente",
        ]

        fill = None
        if score >= 8:
            fill = FILL_GREEN
        elif score >= 6:
            fill = FILL_YELLOW

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            if fill:
                cell.fill = fill

        # Hyperlink en Maps y Landing
        maps_cell = ws.cell(row=row_idx, column=12)
        if maps_cell.value:
            maps_cell.hyperlink = maps_cell.value
            maps_cell.value = "Abrir en Maps"
            maps_cell.font = Font(color="1F6FEB", underline="single")

        landing_cell = ws.cell(row=row_idx, column=13)
        if landing_cell.value:
            # openpyxl acepta rutas locales como hyperlink
            landing_cell.hyperlink = landing_cell.value
            landing_cell.value = "Abrir landing"
            landing_cell.font = Font(color="1F6FEB", underline="single")

    # Dropdown de Estado en la columna N
    dv = DataValidation(
        type="list",
        formula1=f'"{",".join(config.STATUS_OPTIONS)}"',
        allow_blank=True,
        showDropDown=False,  # showDropDown=False = sí muestra (es contraintuitivo en openpyxl)
    )
    dv.error = "Selecciona un estado válido"
    dv.errorTitle = "Estado inválido"
    dv.prompt = "Selecciona el estado del prospecto"
    dv.promptTitle = "Estado"
    last_row = max(2, len(prospects) + 1)
    dv.add(f"N2:N{last_row}")
    ws.add_data_validation(dv)

    # Hoja resumen
    ws2 = wb.create_sheet("Resumen")
    total = len(prospects)
    sin_web = sum(1 for p in prospects if not p.get("website"))
    con_email = sum(1 for p in prospects if p.get("email"))
    con_landing = sum(1 for p in prospects if p.get("landing_path"))
    score_alto = sum(1 for p in prospects if p.get("score", 0) >= 8)

    summary = [
        ("Reporte de prospectos", ""),
        ("Generado", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Agencia", config.AGENCY_NAME),
        ("", ""),
        ("Total prospectos", total),
        ("Sin web", sin_web),
        ("Con email", con_email),
        ("Landings generadas", con_landing),
        ("Score alto (>=8)", score_alto),
    ]
    for i, (k, v) in enumerate(summary, start=1):
        ws2.cell(row=i, column=1, value=k).font = Font(bold=True)
        ws2.cell(row=i, column=2, value=v)
    ws2.column_dimensions["A"].width = 24
    ws2.column_dimensions["B"].width = 32

    # Guardar
    fname = f"prospectos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    path = config.EXCEL_DIR / fname
    wb.save(path)
    logger.info("Excel exportado: %s", path)
    return str(path)
