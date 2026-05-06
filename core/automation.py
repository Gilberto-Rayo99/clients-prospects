"""Pipeline de automatización masiva.

Orquesta:
  Excel → clients.json → (los con landing) → Netlify → URL → mensaje WhatsApp
                       → (los sin landing) → marcar pendiente

Diseño en 2 fases:
  - Fase A: subes el Excel. Importa todos los prospectos a clients.json
    (sin duplicar). Para los que YA tienen landing_html, sube a Netlify y
    arma el link wa.me con mensaje. Para los demás, los marca como
    "Falta landing".
  - Fase B: cargas las landings pendientes (uploader por cliente). Cuando
    ya están todas, vuelves a correr Fase A y se procesan los que
    faltaban.

Esquema de estados que pone el pipeline en clients.json:
  - "Mensaje listo"        → landing publicada, link wa.me listo, falta dar click
  - "Falta landing"        → necesita HTML cargado primero
  - "Sin teléfono"         → no se puede mandar WhatsApp
  - "Error: {razón}"       → algo falló (Netlify, etc.)
"""
from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Callable, Optional

import openpyxl

import config
from core import clients_store, netlify
from core.messaging import render_message, whatsapp_url

logger = logging.getLogger(__name__)


# ============================================================
# Excel → clients.json (importación)
# ============================================================
EXCEL_COLUMN_MAP = {
    "score":       1,
    "name":        2,
    "category":    3,
    "address":     4,
    "phone":       5,
    "email":       6,
    "contact_channel_label": 7,
    "contact_value":         8,
    "rating":     9,
    "reviews_count": 10,
    "tiene_web":  11,
    "maps_url":   12,
    "landing_path": 13,
    "estado":     14,
    "fecha_proximo_contacto": 15,
    "precio_cotizado": 16,
    "notas":      17,
    "prompt_claude": 18,
    "fecha_guardado": 19,
}


def _channel_label_to_key(label: str | None) -> str:
    if not label:
        return "visit"
    label = label.strip().lower()
    return {
        "email":            "email",
        "whatsapp":         "whatsapp",
        "facebook":         "facebook",
        "visita presencial": "visit",
        "visit":            "visit",
    }.get(label, "visit")


def _read_cell(ws, row: int, col: int):
    val = ws.cell(row=row, column=col).value
    if val == "—" or val == "-" or val == "":
        return None
    return val


def _read_hyperlink(ws, row: int, col: int) -> str | None:
    cell = ws.cell(row=row, column=col)
    if cell.hyperlink:
        return cell.hyperlink.target
    val = cell.value
    if isinstance(val, str) and val.startswith("http"):
        return val
    return None


def import_excel_to_store(excel_path: str) -> dict:
    """Lee el Excel y mete en clients.json los que falten.

    Returns:
        {"imported": N, "already_existed": N, "rows_total": N}
    """
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if "Prospectos" not in wb.sheetnames:
        raise ValueError("El Excel no tiene hoja 'Prospectos'")
    ws = wb["Prospectos"]

    imported, existed, rows_total = 0, 0, 0
    for row_idx in range(2, ws.max_row + 1):
        rows_total += 1
        name = _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["name"])
        if not name:
            continue

        # Reconstruir registro de cliente
        phone_raw = _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["phone"])
        rating = _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["rating"])
        reviews = _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["reviews_count"]) or 0
        tiene_web = _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["tiene_web"])
        maps_url = _read_hyperlink(ws, row_idx, EXCEL_COLUMN_MAP["maps_url"])
        notas = _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["notas"]) or ""
        prompt_claude = _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["prompt_claude"])
        estado = _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["estado"]) or "Pendiente"
        precio = _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["precio_cotizado"])
        prox = _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["fecha_proximo_contacto"])

        business = {
            "name": str(name).strip(),
            "category": _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["category"]) or "Otros",
            "address": _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["address"]) or "",
            "phone": str(phone_raw).strip() if phone_raw else None,
            "email": _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["email"]),
            "rating": float(rating) if rating else None,
            "reviews_count": int(reviews) if reviews else 0,
            "website": None if tiene_web in (None, "No") else "https://(tenía web)",
            "place_id": f"excel__{_slug(name)}",  # ID estable derivado del nombre
            "maps_url": maps_url,
            "score": _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["score"]),
            "contact_channel": _channel_label_to_key(
                _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["contact_channel_label"])
            ),
            "contact_value": _read_cell(ws, row_idx, EXCEL_COLUMN_MAP["contact_value"]) or "",
        }

        # ¿Ya existe?
        if clients_store.exists_by_place_id(business["place_id"]):
            existed += 1
            continue

        cli = clients_store.save_from_business(business)
        # Aplicar campos extra (notas, estado, precio)
        updates: dict = {}
        if notas:
            updates["notas"] = str(notas)
        if prompt_claude:
            updates["prompt_claude_excel"] = str(prompt_claude)
        if estado and estado != "Pendiente":
            updates["estado"] = str(estado)
        if precio:
            try:
                updates["precio_cotizado"] = float(precio)
            except (TypeError, ValueError):
                pass
        if prox:
            updates["fecha_proximo_contacto"] = str(prox)
        if updates:
            clients_store.update(cli["id"], **updates)
        imported += 1

    return {"imported": imported, "already_existed": existed, "rows_total": rows_total}


def _slug(s: str) -> str:
    import re
    out = re.sub(r"[^a-z0-9]+", "-", str(s).lower()).strip("-")
    return out or "cliente"


# ============================================================
# Fase A — procesar lote
# ============================================================
def run_phase_a(
    client_ids: list[str],
    template_key: str = "inicial_sin_web",
    progress_cb: Optional[Callable[[int, int, str], None]] = None,
) -> list[dict]:
    """Procesa cada cliente: si tiene landing_html, sube a Netlify y arma mensaje.

    Args:
        client_ids: lista de IDs en clients.json a procesar
        template_key: plantilla de mensaje WhatsApp
        progress_cb: callback(idx, total, status_message) para reportar

    Returns:
        Lista de dicts con resultado por cliente:
          {"id", "name", "ok", "url", "wa_url", "estado", "error"}
    """
    results = []
    total = len(client_ids)

    for i, cid in enumerate(client_ids, start=1):
        cli = clients_store.get(cid)
        if not cli:
            results.append({
                "id": cid, "name": "(no existe)", "ok": False,
                "estado": "Error: cliente no existe", "error": "not_found",
            })
            continue

        name = cli.get("name", "")
        if progress_cb:
            progress_cb(i, total, f"Procesando: {name}")

        # 1) Sin teléfono → no se puede mandar WhatsApp
        if not cli.get("phone"):
            clients_store.update(cid, estado="Sin teléfono")
            results.append({
                "id": cid, "name": name, "ok": False,
                "estado": "Sin teléfono", "error": "no_phone",
            })
            continue

        # 2) Sin landing → marcar pendiente
        if not cli.get("landing_html"):
            clients_store.update(cid, estado="Falta landing")
            results.append({
                "id": cid, "name": name, "ok": False,
                "estado": "Falta landing", "error": "no_landing",
            })
            continue

        # 3) Subir a Netlify (reusar site_id si existe)
        try:
            res = netlify.publish_html(
                cli["landing_html"], name, site_id=cli.get("netlify_site_id"),
            )
        except netlify.NetlifyError as e:
            logger.exception("Netlify falló para %s", name)
            clients_store.update(cid, estado=f"Error: Netlify ({e})")
            results.append({
                "id": cid, "name": name, "ok": False,
                "estado": f"Error: Netlify", "error": str(e),
            })
            continue

        # 4) Validar URL viva
        if not netlify.url_is_alive(res["url"]):
            logger.warning("URL %s no responde — sigo de todos modos", res["url"])

        # 5) Generar mensaje WhatsApp + link wa.me
        rendered = render_message(template_key, cli, res["url"])
        wa = whatsapp_url(cli, rendered)

        # 6) Persistir — no actualizar deploy_at si fue skipped (sin deploy real)
        update_fields = {
            "netlify_url": res["url"],
            "netlify_site_id": res["site_id"],
            "estado": "Mensaje listo",
        }
        if not res.get("skipped"):
            update_fields["netlify_deploy_at"] = datetime.now().isoformat(timespec="seconds")
        clients_store.update(cid, **update_fields)

        results.append({
            "id": cid, "name": name, "ok": True,
            "url": res["url"], "wa_url": wa,
            "estado": "Mensaje listo", "error": None,
        })

    return results


# ============================================================
# Excel actualizado regenerado desde clients.json
# ============================================================
def regenerate_excel_from_store(only_ids: list[str] | None = None) -> str:
    """Regenera Excel con datos actuales de clients.json.

    Args:
        only_ids: si se pasa, solo exporta esos. Si es None, exporta todos.

    Returns:
        Ruta del Excel generado.
    """
    from export.excel import export_to_excel

    all_clients = clients_store.list_all()
    if only_ids is not None:
        all_clients = [c for c in all_clients if c["id"] in set(only_ids)]

    # Añadir el campo netlify_url al landing_path para que aparezca como link en el Excel
    enriched = []
    for c in all_clients:
        c2 = dict(c)
        if c.get("netlify_url"):
            c2["landing_path"] = c["netlify_url"]
        enriched.append(c2)

    path = export_to_excel(enriched)
    # Renombrar con sufijo _automatizado
    p = Path(path)
    new_name = p.with_name(p.stem.replace("prospectos_", "prospectos_automatizado_") + p.suffix)
    p.rename(new_name)
    logger.info("Excel automatizado: %s", new_name)
    return str(new_name)
